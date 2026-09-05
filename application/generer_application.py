# -*- coding: utf-8 -*-
"""
Construit l'application « Alternance 974 » en injectant dans le modele HTML
les structures issues des deux fichiers de prospection.

    python3 generer_application.py

Produit : application/Alternance_974.html — un fichier unique, autonome,
qui s'ouvre par double-clic sur Mac comme sur PC.
"""
import io
import json
import os
import re
import sys
import unicodedata

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl manquant : python3 -m pip install openpyxl")

ICI = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.dirname(ICI)
PROSPECTION = os.path.join(RACINE, "prospection")
MODELE = os.path.join(ICI, "modele_application.html")
SORTIE = os.path.join(ICI, "Alternance_974.html")

FICHIERS = [
    ("prive",  os.path.join(PROSPECTION, "Prospection_Logistique_Nord-Est_Reunion.xlsx")),
    ("public", os.path.join(PROSPECTION, "Prospection_Public_Sante_Nord-Est_Reunion.xlsx")),
]

# Colonnes du classeur (voir prospection/commun_prospection.py)
C_NOM, C_TYPE, C_ADR, C_COM = 2, 3, 4, 6
C_TEL, C_MAIL, C_SITE, C_POSTE, C_PRIO, C_NOTES = 13, 14, 15, 16, 17, 22

PLACEHOLDERS = {"A VERIFIER", "A COMPLETER", "", "None"}

# Poste du fichier de prospection -> metier de l'application
METIERS = {
    "opérateur logistique": "prepa",
    "magasinier": "magasinier",
    "agent de quai": "quai",
    "agent de transit": "transit",
    "employé logistique": "elog",
    "magasin communal": "magasinier",
    "économat": "econome",
    "restauration collective": "econome",
    "logistique pharmaceutique": "pharma",
    "parc matériel": "magasinier",
}

COMMUNES = {
    "saint-denis": "saint-denis",
    "sainte-clotilde": "sainte-clotilde",
    "saint-denis (hauteurs)": "saint-denis-hauts",
    "sainte-marie": "sainte-marie",
    "sainte-suzanne": "sainte-suzanne",
    "saint-andre": "saint-andre",
    "salazie": "salazie",
    "bras-panon": "bras-panon",
    "saint-benoit": "saint-benoit",
    "sainte-anne": "sainte-anne",
    "la plaine-des-palmistes": "plaine-palmistes",
    "sainte-rose": "sainte-rose",
    "bassin nord-est": "multi",
    "saint-paul": "hors",
}


def propre(v):
    t = str(v).strip() if v is not None else ""
    return "" if t in PLACEHOLDERS else t


def sans_accent(t):
    return "".join(c for c in unicodedata.normalize("NFD", t)
                   if unicodedata.category(c) != "Mn")


def id_commune(libelle):
    """'Sainte-Clotilde (Saint-Denis)' -> 'sainte-clotilde'."""
    t = sans_accent(str(libelle or "")).lower().strip()
    if t in COMMUNES:
        return COMMUNES[t]
    base = re.sub(r"\(.*?\)", "", t).strip()
    if base in COMMUNES:
        return COMMUNES[base]
    for cle, val in COMMUNES.items():          # ex. "sainte-clotilde (saint-denis)"
        if base.startswith(cle):
            return val
    return ""


def id_metier(poste):
    t = str(poste or "").lower()
    for cle, val in METIERS.items():
        if cle in t:
            return val
    return "prepa"


def lire(fichier, secteur):
    wb = load_workbook(fichier, data_only=True)
    lignes = []
    for onglet in wb.sheetnames:
        if not re.match(r"^0[1-7] ", onglet):
            continue
        ws = wb[onglet]
        for r in range(5, ws.max_row + 1):
            nom = propre(ws.cell(row=r, column=C_NOM).value)
            if not nom or nom.startswith(("À prospecter", "A prospecter")):
                continue
            type_structure = propre(ws.cell(row=r, column=C_TYPE).value)
            if "à lister" in type_structure.lower() or "à défricher" in type_structure.lower():
                continue          # lignes methode, pas des structures reelles
            commune = id_commune(ws.cell(row=r, column=C_COM).value)
            if not commune:
                print("  ! commune non reconnue, ligne ignorée :", nom,
                      "/", ws.cell(row=r, column=C_COM).value)
                continue
            poste = propre(ws.cell(row=r, column=C_POSTE).value)
            adresse = propre(ws.cell(row=r, column=C_ADR).value)
            notes = propre(ws.cell(row=r, column=C_NOTES).value)
            lignes.append({
                "id": "p" + str(len(lignes) + 1) + secteur[0],
                "structure": nom,
                "type": type_structure,
                "secteur": secteur,
                "commune": commune,
                "metier": id_metier(poste),
                "poste": poste,
                "postes": 1,
                "contact": "",
                "tel": propre(ws.cell(row=r, column=C_TEL).value),
                "email": propre(ws.cell(row=r, column=C_MAIL).value),
                "site": propre(ws.cell(row=r, column=C_SITE).value),
                "priorite": propre(ws.cell(row=r, column=C_PRIO).value) or "P2",
                "statut": "A contacter",
                "notes": " — ".join(x for x in (adresse, notes) if x),
                "source": onglet,
            })
    return lignes


def main():
    toutes = []
    for secteur, fichier in FICHIERS:
        if not os.path.exists(fichier):
            sys.exit("Fichier de prospection introuvable : %s\n"
                     "Générez-le d'abord depuis le dossier prospection/." % fichier)
        lues = lire(fichier, secteur)
        print("%-7s %2d structures  (%s)" % (secteur, len(lues), os.path.basename(fichier)))
        toutes.extend(lues)

    # identifiants uniques et stables
    for i, o in enumerate(toutes, start=1):
        o["id"] = "ref%03d" % i

    modele = io.open(MODELE, encoding="utf-8").read()
    if "/*__DONNEES__*/" not in modele:
        sys.exit("Le modèle ne contient plus le repère /*__DONNEES__*/.")
    donnees = json.dumps(toutes, ensure_ascii=False, indent=1)
    html = modele.replace("/*__DONNEES__*/[]", donnees, 1)

    io.open(SORTIE, "w", encoding="utf-8").write(html)
    print("\nOK -> %s" % SORTIE)
    print("     %d structures pré-chargées, %d Ko" % (len(toutes), len(html.encode()) // 1024))


if __name__ == "__main__":
    main()
