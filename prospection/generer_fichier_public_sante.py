# -*- coding: utf-8 -*-
"""
Genere le fichier de prospection APPRENTISSAGE - structures publiques et sante
Bassin Nord-Est de La Reunion (Saint-Denis -> Sainte-Rose, hauteurs incluses).

Cible : toute structure NON marchande ou de sante disposant d'un magasin, d'un
economat, d'une cuisine centrale, d'une pharmacie ou d'un parc materiel, et donc
capable d'accueillir un(e) apprenti(e) operateur logistique de 16 a 25 ans :
collectivites (communes, CCAS, intercommunalites, Departement, Region), ecoles,
lycees et campus, hopitaux, cliniques, dialyse, EHPAD et medico-social.

IMPORTANT - INTEGRITE DES DONNEES
Les noms des structures sont publics et verifiables. Les adresses proviennent de
sources publiques et restent a confirmer. Les champs SIRET / Code APE /
Responsable / Telephone / Email ne sont renseignes QUE lorsqu'ils ont ete
verifies. Sinon : "A VERIFIER" + lien vers l'Annuaire des Entreprises.
Aucune donnee n'a ete inventee.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from commun_prospection import (
    AC, AV, BLEU, BORD, F_CORPS, F_ENTETE, F_GRAS, F_H2, F_LIEN, F_STITRE,
    F_TITRE, FILL_ALT, FILL_ENTETE, FILL_JAUNE, L, P_MAGAS, P_OPLOG,
    colonnes, entete_feuille, feuille_commune, onglet_pilotage,
)

COLS = colonnes("Type de structure")

# ---------------------------------------------------------------- Postes vises
P_MAGCOM = "Magasinier / Magasin communal"
P_ECON   = "Agent d'économat / Réception-stocks"
P_RESTO  = "Logistique de restauration collective"
P_PHARMA = "Logistique pharmaceutique / PUI"
P_ATEL   = "Magasinier atelier / Parc matériel"

# Rappels d'interlocuteur, repetes dans les notes car c'est LA cle du public
I_COMMUNE = ("Interlocuteurs : DGS + DRH + directeur des services techniques "
             "(c'est lui qui a le magasin).")
I_LYCEE = ("Interlocuteur : l'ADJOINT GESTIONNAIRE (intendant). C'est lui qui gère "
           "le magasin et la restauration ; le chef d'établissement signe.")
I_HOPITAL = ("Interlocuteurs : direction des achats et de la logistique + DRH. "
             "Le cadre du magasin général est le prescripteur.")

# ---------------------------------------------------------------- Donnees
COMMUNES = [
    L("VILLE DE SAINT-DENIS", "Commune (fonction publique territoriale)",
      "2 rue de Paris – Hôtel de Ville", "97400", "Saint-Denis", P_MAGCOM, "P1",
      "Source annuaire",
      "La plus grosse collectivité de l'île : magasin général, cuisines scolaires, "
      "parc matériel des fêtes et cérémonies, archives, parc automobile. Plusieurs "
      "postes d'apprentis logistique possibles la même année. " + I_COMMUNE),
    L("CCAS DE SAINT-DENIS", "CCAS (établissement public communal)",
      AV, "97400", "Saint-Denis", P_RESTO, "P2", "Source annuaire",
      "Portage de repas à domicile, aide alimentaire, magasin de fournitures : "
      "flux quotidiens et gestion de stock du froid. Interlocuteur : directeur du CCAS."),
    L("VILLE DE SAINTE-MARIE", "Commune (fonction publique territoriale)",
      AV, "97438", "Sainte-Marie", P_MAGCOM, "P1", "A vérifier",
      "Commune en croissance, zone aéroportuaire sur son territoire. Magasin communal, "
      "restauration scolaire, services techniques. " + I_COMMUNE),
    L("CCAS DE SAINTE-MARIE", "CCAS (établissement public communal)",
      AV, "97438", "Sainte-Marie", P_RESTO, "P3", "A vérifier",
      "Portage de repas et aide sociale : vérifier le volume avant de proposer un apprenti."),
    L("VILLE DE SAINTE-SUZANNE", "Commune (fonction publique territoriale)",
      AV, "97441", "Sainte-Suzanne", P_MAGCOM, "P2", "A vérifier",
      "Commune de taille moyenne : un seul magasin, mais un tuteur facile à identifier. "
      + I_COMMUNE),
    L("VILLE DE SAINT-ANDRÉ", "Commune (fonction publique territoriale)",
      AV, "97440", "Saint-André", P_MAGCOM, "P1", "A vérifier",
      "2e ville de l'Est. Restauration scolaire importante, services techniques étoffés, "
      "magasin communal. Cible prioritaire du bassin Est. " + I_COMMUNE),
    L("CCAS DE SAINT-ANDRÉ", "CCAS (établissement public communal)",
      AV, "97440", "Saint-André", P_RESTO, "P2", "A vérifier",
      "Action sociale et portage de repas sur un bassin de population important."),
    L("VILLE DE BRAS-PANON", "Commune (fonction publique territoriale)",
      AV, "97412", "Bras-Panon", P_MAGCOM, "P2", "A vérifier",
      "Petite commune : le maire et le DGS décident vite. Magasin et parc matériel. "
      + I_COMMUNE),
    L("VILLE DE SAINT-BENOÎT", "Commune (fonction publique territoriale)",
      AV, "97470", "Saint-Benoît", P_MAGCOM, "P1", "A vérifier",
      "Ville-centre de l'Est, siège de la CIREST. Services techniques, restauration "
      "scolaire, magasin. " + I_COMMUNE),
    L("VILLE DE LA PLAINE-DES-PALMISTES", "Commune (fonction publique territoriale)",
      AV, "97431", "La Plaine-des-Palmistes", P_MAGCOM, "P3", "A vérifier",
      "Commune des hauteurs, effectifs réduits : viser un poste polyvalent magasin + "
      "régie. Vérifier qu'un tuteur est disponible."),
    L("VILLE DE SAINTE-ROSE", "Commune (fonction publique territoriale)",
      AV, "97439", "Sainte-Rose", P_MAGCOM, "P3", "A vérifier",
      "Commune la plus éloignée du bassin : c'est souvent le SEUL employeur structuré "
      "de la commune. Argument fort pour un jeune du secteur qui ne peut pas se déplacer."),
    L("VILLE DE SALAZIE", "Commune (fonction publique territoriale)",
      AV, "97433", "Salazie", P_MAGCOM, "P3", "A vérifier",
      "Cirque de Salazie : mobilité très contrainte pour les jeunes. La mairie est le "
      "principal employeur local. Même argument que Sainte-Rose."),
]

INTERCOS = [
    L("CIREST – COMMUNAUTÉ INTERCOMMUNALE RÉUNION EST", "Intercommunalité (EPCI)",
      AV, "97437", "Saint-Benoît", P_MAGAS, "P1", "Source annuaire",
      "Regroupe 6 communes : Bras-Panon, La Plaine-des-Palmistes, Saint-André, "
      "Saint-Benoît, Sainte-Rose et Salazie. SIREN relevé sur l'Annuaire des "
      "Entreprises : 249740093 (le SIRET du siège reste à récupérer). Compétences "
      "déchets, transport, eau : magasins, ateliers et flux de matériel.",
      site="https://www.cirest.fr/"),
    L("CINOR – COMMUNAUTÉ INTERCOMMUNALE DU NORD DE LA RÉUNION", "Intercommunalité (EPCI)",
      AV, "97400", "Saint-Denis", P_MAGAS, "P1", "Source annuaire",
      "Regroupe Saint-Denis, Sainte-Marie et Sainte-Suzanne. Compétences déchets, "
      "transports, développement économique : ateliers, magasins, parc de bacs et "
      "de conteneurs à gérer."),
    L("SYDNE – SYNDICAT MIXTE DE TRAITEMENT DES DÉCHETS", "Syndicat mixte",
      AV, "97400", "Saint-Denis", P_OPLOG, "P2", "A vérifier",
      "Syndicat de traitement des déchets porté par la CINOR et la CIREST. "
      "Flux, pesée, quais de transfert : environnement logistique réel. "
      "Confirmer le siège et la commune avant l'appel."),
    L("SIDELEC RÉUNION", "Syndicat intercommunal d'électricité",
      AV, "97400", "Saint-Denis", P_ATEL, "P3", "A vérifier",
      "Syndicat d'électrification : magasin de matériel électrique et parc. "
      "Vérifier l'existence d'un magasin propre avant de prospecter."),
]

DEPARTEMENT_REGION_ETAT = [
    L("CONSEIL DÉPARTEMENTAL DE LA RÉUNION", "Département (collectivité)",
      "2 rue de la Source – Hôtel du Département", "97400", "Saint-Denis", P_MAGAS, "P1",
      "Source annuaire",
      "Très gros employeur territorial. Gère les COLLÈGES (restauration et matériel), "
      "les routes départementales, l'action sociale : magasins, économats et parcs "
      "sur tout le territoire, y compris à l'Est. Passer par la DRH ET la direction "
      "de la logistique / des moyens généraux."),
    L("RÉGION RÉUNION", "Région (collectivité)",
      "Avenue René Cassin – Moufia, Hôtel de Région Pierre Lagourgue", "97490",
      "Sainte-Clotilde (Saint-Denis)", P_MAGAS, "P1", "Source annuaire",
      "Gère les LYCÉES (restauration, équipement, maintenance) et les grandes "
      "infrastructures. Deux portes d'entrée : les moyens généraux de la Région, "
      "et chaque lycée pris individuellement (voir onglet 04)."),
    L("RECTORAT DE L'ACADÉMIE DE LA RÉUNION", "État (éducation nationale)",
      "24 avenue Georges Brassens – Le Moufia", "97490", "Sainte-Clotilde (Saint-Denis)",
      P_ECON, "P2", "Source annuaire",
      "Fonction publique d'État : apprentissage possible. Services logistiques, "
      "reprographie, magasin. Ouvre surtout la porte des lycées et collèges du bassin."),
    L("UNIVERSITÉ DE LA RÉUNION", "Établissement public d'enseignement supérieur",
      "15 avenue René Cassin – Campus du Moufia", "97490", "Sainte-Clotilde (Saint-Denis)",
      P_MAGAS, "P2", "Source annuaire",
      "Magasin central, laboratoires, reprographie, logistique de campus. "
      "Structure habituée à l'alternance : dossier administratif bien rodé."),
    L("CROUS DE LA RÉUNION", "Établissement public (vie étudiante)",
      AV, "97490", "Sainte-Clotilde (Saint-Denis)", P_RESTO, "P2", "A vérifier",
      "Restauration universitaire et résidences : réception de denrées, gestion de "
      "stocks, approvisionnement de plusieurs points de vente. Vrai métier logistique."),
    L("SDIS 974 – SERVICE D'INCENDIE ET DE SECOURS", "Établissement public territorial",
      AV, "97400", "Saint-Denis", P_MAGAS, "P1", "A vérifier",
      "Magasin d'habillement, magasin de matériel de secours, pharmacie, parc de "
      "véhicules, approvisionnement de tous les centres de secours du département. "
      "L'un des plus beaux terrains logistiques du public. Confirmer l'adresse de "
      "l'état-major et le service en charge de la logistique."),
]

ECOLES = [
    L("LYCÉE LISLET GEOFFROY", "Lycée (EPLE)",
      AV, "97490", "Sainte-Clotilde (Saint-Denis)", P_ECON, "P1", "Source annuaire",
      "Magasin d'intendance et restauration scolaire. " + I_LYCEE),
    L("LYCÉE SARDA GARRIGA", "Lycée (EPLE)",
      AV, "97440", "Saint-André", P_ECON, "P1", "Source annuaire",
      "Gros établissement de l'Est : réception de denrées quotidienne, magasin, "
      "gestion des manuels et des équipements. " + I_LYCEE),
    L("LYCÉE MAHATMA GANDHI", "Lycée (EPLE)",
      AV, "97440", "Saint-André", P_ECON, "P2", "Source annuaire",
      "Restauration scolaire et magasin. " + I_LYCEE),
    L("LYCÉE AMIRAL PIERRE BOUVET", "Lycée (EPLE)",
      AV, "97470", "Saint-Benoît", P_ECON, "P1", "Source annuaire",
      "Établissement de référence de Saint-Benoît. " + I_LYCEE),
    L("LYCÉE DE BRAS-FUSIL", "Lycée professionnel (EPLE)",
      AV, "97470", "Saint-Benoît", P_ECON, "P1", "Source annuaire",
      "Lycée professionnel : la culture de l'alternance y est déjà installée, "
      "l'accueil d'un apprenti y est plus naturel. " + I_LYCEE),
    L("AUTRES LYCÉES ET COLLÈGES DU BASSIN", "Établissements scolaires – à lister",
      "Annuaire de l'éducation (data.education.gouv.fr) filtré sur les communes du bassin",
      "97400", "Bassin Nord-Est", P_ECON, "P2", "A vérifier",
      "Chaque collège et lycée a un adjoint gestionnaire, une demi-pension et un magasin. "
      "C'est un gisement de postes très peu prospecté par les CFA. "
      "Méthode : annuaire de l'éducation, filtre commune, puis appel direct à l'intendance."),
]

HOPITAL_PUBLIC = [
    L("CHU DE LA RÉUNION – SITE FÉLIX GUYON", "CHU (fonction publique hospitalière)",
      "Allée des Topazes – Bellepierre", "97400", "Saint-Denis", P_PHARMA, "P1",
      "Source annuaire",
      "LE plus gros employeur logistique du Nord : magasin général, plateforme de "
      "distribution interne, pharmacie à usage intérieur, blanchisserie, restauration, "
      "brancardage. Le CHU dispose d'une Direction des Achats et de la Logistique (Nord). "
      + I_HOPITAL,
      site="https://www.chu-reunion.fr/"),
    L("GROUPE HOSPITALIER EST RÉUNION (GHER)", "Hôpital public (FPH)",
      "30 RN3 – ZAC Madeleine, Bras-Fusil", "97470", "Saint-Benoît", P_PHARMA, "P1",
      "Source annuaire",
      "Hôpital de référence de l'Est : magasin général, pharmacie, restauration. "
      "ATTENTION : cette structure figure aussi dans le fichier logistique privée — "
      "ne passez pas deux appels séparés. " + I_HOPITAL),
    L("EPSMR – ÉTABLISSEMENT PUBLIC DE SANTÉ MENTALE", "Hôpital public (FPH) – hors bassin",
      AV, "97460", "Saint-Paul", P_MAGAS, "P3", "A vérifier",
      "HORS BASSIN Nord-Est : à ne proposer qu'à un candidat mobile et véhiculé. "
      "Confirmer le site avant tout contact."),
]

SANTE_PRIVEE = [
    L("CLINIQUE SAINTE-CLOTILDE (GROUPE CLINIFUTUR)", "Clinique privée MCO",
      "127 route de Bois de Nèfles – Sainte-Clotilde", "97490",
      "Sainte-Clotilde (Saint-Denis)", P_PHARMA, "P1", "Source annuaire",
      "Environ 650 salariés et 100 médecins : la plus grande clinique de l'île et le "
      "plus gros plateau chirurgical privé du Nord-Est. Pharmacie à usage intérieur, "
      "magasin de dispositifs médicaux, logistique hôtelière et restauration. "
      "Interlocuteurs : DRH + pharmacien gérant + responsable des services économiques.",
      site="https://www.clinifutur.net/"),
    L("GROUPE DE SANTÉ CLINIFUTUR – direction", "Groupe de santé privé (siège)",
      AV, "97490", "Sainte-Clotilde (Saint-Denis)", P_MAGAS, "P1", "Source annuaire",
      "Le groupe exploite plusieurs cliniques à La Réunion et à Mayotte "
      "(Sainte-Clotilde, La Paix, Les Oliviers, Robert Debré, Saint-Joseph, "
      "Saint-Vincent) ainsi que les centres de dialyse SODIA. Un accord au niveau du "
      "groupe peut ouvrir plusieurs sites d'un coup : c'est le meilleur rapport "
      "effort / résultat de cet onglet."),
    L("SODIA RÉUNION (dialyse – groupe Clinifutur)", "Dialyse / Santé privée",
      AV, "97490", "Sainte-Clotilde (Saint-Denis)", P_OPLOG, "P2", "A vérifier",
      "Centres de dialyse : consommables, poches, filtres, livraisons sur plusieurs "
      "sites. Logistique du soin très normée. Confirmer les sites du bassin Nord-Est."),
    L("AUTRES CLINIQUES DU GROUPE (La Paix, Les Oliviers, Robert Debré, Saint-Joseph, Saint-Vincent)",
      "Cliniques privées – à lister",
      AV, "97400", "À vérifier", P_MAGAS, "P3", "A vérifier",
      "Ces cliniques existent, mais TOUTES ne sont pas dans le bassin Nord-Est. "
      "Vérifiez la commune de chacune avant de les intégrer à votre tournée."),
]

MEDICO_SOCIAL = [
    L("AURAR – UNITÉ DE SAINT-BENOÎT", "Dialyse / Association de santé",
      "1 rue des Aubépines", "97470", "Saint-Benoît", P_OPLOG, "P1", "Source annuaire",
      "Centre lourd, dialyse médicalisée et auto-dialyse. Téléphone relevé en source "
      "publique, à reconfirmer. Approvisionnement en consommables et gestion de stock "
      "sur site : poste d'apprenti tout à fait tenable.",
      tel="0262 98 98 98", site="https://aurar.fr/nous-trouver/"),
    L("AURAR – AUTRES UNITÉS DU NORD-EST", "Dialyse / Association de santé",
      AV, "97490", "Sainte-Clotilde (Saint-Denis)", P_OPLOG, "P2", "Source annuaire",
      "L'AURAR est présente dans l'Est, le Nord, l'Ouest et le Sud. Une seule "
      "négociation avec le siège peut couvrir plusieurs unités du bassin."),
    L("EHPAD DU BASSIN NORD-EST", "EHPAD – à lister",
      "Annuaire sante.fr / FINESS, filtre communes du bassin", "97400", "Bassin Nord-Est",
      P_RESTO, "P2", "A vérifier",
      "Chaque EHPAD a une cuisine, une lingerie, un magasin de protections et de "
      "matériel médical. Structures petites mais nombreuses : un apprenti par EHPAD, "
      "et le directeur décide seul. Méthode : sante.fr, filtre EHPAD + commune."),
    L("ESAT ET ÉTABLISSEMENTS MÉDICO-SOCIAUX DU BASSIN", "ESAT / IME / Foyers – à lister",
      AV, "97400", "Bassin Nord-Est", P_OPLOG, "P2", "A vérifier",
      "Les ESAT ont souvent des ateliers de conditionnement, de tri et de logistique : "
      "l'apprenti y encadre et organise le flux. Double intérêt : poste logistique réel "
      "et forte adhésion des directions à la démarche d'insertion."),
]

ONGLETS = [
    ("01 Communes & CCAS", COMMUNES,
     "Les 9 communes du bassin et leurs CCAS – magasin communal, restauration scolaire, parc matériel"),
    ("02 Intercommunalités", INTERCOS,
     "CIREST, CINOR et syndicats – déchets, transports, ateliers et magasins"),
    ("03 Département, Région & État", DEPARTEMENT_REGION_ETAT,
     "Département, Région, Rectorat, Université, CROUS, SDIS – les gros employeurs publics"),
    ("04 Lycées & collèges", ECOLES,
     "EPLE du bassin – magasin d'intendance et restauration scolaire (voir l'adjoint gestionnaire)"),
    ("05 Hôpital public", HOPITAL_PUBLIC,
     "CHU site Félix Guyon, GHER – magasin général, pharmacie à usage intérieur, plateforme"),
    ("06 Cliniques & santé privée", SANTE_PRIVEE,
     "Groupe Clinifutur et cliniques du Nord-Est – PUI, dispositifs médicaux, logistique hôtelière"),
    ("07 EHPAD & médico-social", MEDICO_SOCIAL,
     "AURAR, EHPAD, ESAT – petites structures, décision rapide, un apprenti chacune"),
]

# ---------------------------------------------------------------- Workbook
wb = Workbook()
wb.remove(wb.active)

# ---- Onglet 00 : mode d'emploi
ws = wb.create_sheet("00 Mode d'emploi")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 40
ws.column_dimensions["C"].width = 100
ws["B2"] = "PROSPECTION APPRENTISSAGE – SECTEUR PUBLIC, ÉCOLES ET SANTÉ"
ws["B2"].font = Font(name="Calibri", size=18, bold=True, color=BLEU)
ws["B3"] = "Bassin Nord-Est de La Réunion : de Saint-Denis à Sainte-Rose, hauteurs comprises"
ws["B3"].font = F_STITRE

blocs = [
    ("À QUOI SERT CE FICHIER", ""),
    ("Objectif", "Identifier les structures publiques, scolaires et de santé du bassin Nord-Est "
                 "qui disposent d'un magasin, d'un économat, d'une cuisine centrale, d'une "
                 "pharmacie ou d'un parc matériel — donc capables d'accueillir un(e) apprenti(e) "
                 "opérateur logistique de 16 à 25 ans."),
    ("Complémentaire du fichier privé", "Ce classeur est le pendant du fichier "
                 "« Prospection_Logistique_Nord-Est_Reunion.xlsx » (entreprises privées). "
                 "Même structure, mêmes colonnes, même méthode. Une seule structure figure dans "
                 "les deux : le GHER à Saint-Benoît — ne l'appelez pas deux fois."),
    ("Pourquoi ce gisement est sous-exploité", "Les CFA prospectent le privé et oublient le public. "
                 "Or une mairie, un lycée ou un hôpital ont un magasin, des livraisons "
                 "quotidiennes et des agents proches de la retraite à remplacer. "
                 "La concurrence y est bien plus faible."),
    ("", ""),
    ("⚠️ À LIRE AVANT LE PREMIER APPEL", ""),
    ("Le public a un calendrier inversé", "Dans le privé, on prospecte pour la rentrée. Dans le "
                 "public territorial, le financement de la formation par le CNFPT dépend d'un "
                 "RECENSEMENT annuel qui se fait en début d'année civile. Une collectivité qui "
                 "ne s'est pas recensée ne pourra pas être financée. Voir l'onglet "
                 "« 09 Apprentissage dans le public » AVANT de décrocher le téléphone."),
    ("Le bon interlocuteur n'est pas le maire", "Dans une commune : le DGS, la DRH et le directeur "
                 "des services techniques. Dans un lycée : l'adjoint gestionnaire. Dans un "
                 "hôpital : la direction des achats et de la logistique. Se tromper de porte "
                 "fait perdre trois semaines."),
    ("Fiabilité des données", "Les NOMS des structures sont publics et fiables. Les ADRESSES sont "
                 "à confirmer (colonne « Fiabilité donnée »). SIRET, code APE et responsable "
                 "portent « A VERIFIER » tant qu'ils n'ont pas été confirmés : aucune donnée "
                 "n'a été inventée."),
    ("", ""),
    ("REMPLISSAGE AUTOMATIQUE – 2 MINUTES", ""),
    ("Le script d'auto-complétion", "Le fichier « completer_via_annuaire.py » fonctionne aussi sur "
                 "ce classeur : il interroge l'API officielle de l'INSEE et remplit SIRET, code APE, "
                 "effectif et responsable. Les collectivités et les établissements publics sont "
                 "bien présents dans cette base."),
    ("Comment le lancer", "python3 completer_via_annuaire.py "
                 "Prospection_Public_Sante_Nord-Est_Reunion.xlsx — voir LISEZ-MOI.md."),
    ("", ""),
    ("LES ONGLETS", ""),
    ("01 à 07", "Un onglet par famille de structures. Les onglets 01 et 04 sont les plus gros "
                 "gisements en nombre de postes ; les onglets 05 et 06 les plus gros en volume "
                 "logistique par structure."),
    ("08 Pilotage", "Compteurs automatiques par onglet et par statut."),
    ("09 Apprentissage dans le public", "Qui finance quoi, le calendrier CNFPT, les pièces à "
                 "déposer et les délais. C'est l'onglet qui fait la différence face à un DRH."),
    ("10 Script d'appel (public)", "Trame d'appel adaptée : on ne parle pas à une mairie comme "
                 "à un transporteur."),
    ("11 Sources & mise à jour", "D'où viennent les données et où les actualiser."),
]
r = 5
for titre, txt in blocs:
    if titre and not txt:
        ws.cell(row=r, column=2, value=titre).font = F_H2
        ws.cell(row=r, column=2).fill = FILL_JAUNE
        ws.cell(row=r, column=3).fill = FILL_JAUNE
        r += 1
        continue
    if not titre and not txt:
        r += 1
        continue
    a = ws.cell(row=r, column=2, value=titre); a.font = F_GRAS
    a.alignment = Alignment(vertical="top", wrap_text=True)
    b = ws.cell(row=r, column=3, value=txt); b.font = F_CORPS
    b.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = max(15, 13 * (len(txt) // 95 + 1))
    r += 1
ws.sheet_view.showGridLines = False

# ---- Onglets structures
compteurs = []
for nom_onglet, lignes, st in ONGLETS:
    _, n = feuille_commune(wb, nom_onglet, lignes, st,
                           prefixe_titre="PROSPECTION PUBLIC & SANTÉ", cols=COLS)
    compteurs.append((nom_onglet, n))

# ---- Onglet 08 : pilotage
onglet_pilotage(
    wb, compteurs,
    "PILOTAGE DE LA CAMPAGNE – PUBLIC & SANTÉ",
    "Compteurs automatiques – se mettent à jour quand vous remplissez les onglets 01 à 07",
    "TOTAL SECTEUR PUBLIC & SANTÉ",
    [
        ("Structures à rencontrer / semaine", "5"),
        ("Objectif recensement CNFPT", "toutes les collectivités P1 avant la campagne de janvier"),
        ("Offres d'apprentissage à collecter", "à définir"),
        ("Rappel", "dans le public, compter 3 mois entre l'accord oral et la signature"),
    ],
    libelle_volume="VOLUME PAR FAMILLE DE STRUCTURES", libelle_unite="Structures",
)

# ---- Onglet 09 : apprentissage dans le public
ws = wb.create_sheet("09 Apprentissage dans le public")
entete_feuille(ws, "L'APPRENTISSAGE DANS LE SECTEUR PUBLIC",
               "Ce qui change par rapport à une entreprise privée – à maîtriser avant le premier rendez-vous", 2)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 108
public = [
    ("H", "1. TROIS FONCTIONS PUBLIQUES, TROIS RÈGLES", ""),
    ("T", "Territoriale (FPT)", "Communes, CCAS, intercommunalités, Département, Région. "
          "C'est le CNFPT qui prend en charge les frais de formation."),
    ("T", "Hospitalière (FPH)", "CHU, GHER. Le CNFPT n'intervient PAS : le financement relève de "
          "l'établissement, avec l'appui possible de l'ANFH. À confirmer établissement par établissement."),
    ("T", "État (FPE)", "Rectorat, lycées et collèges, université. Financement porté par "
          "l'administration employeuse. À confirmer au cas par cas."),
    ("T", "Le privé de la santé", "Cliniques, dialyse, EHPAD associatifs : ce sont des employeurs "
          "PRIVÉS. On y applique les règles habituelles (OPCO, aides à l'embauche) — c'est "
          "souvent le circuit le plus rapide de tout ce fichier."),
    ("H", "2. CE QUE FINANCE LE CNFPT (COLLECTIVITÉS)", ""),
    ("T", "Le principe", "Le CNFPT prend en charge les frais de formation des apprentis des "
          "employeurs territoriaux. Pour 2026, il annonce le financement de 5 000 nouveaux contrats "
          "au niveau national."),
    ("T", "Les niveaux financés", "Niveaux 3 à 5, c'est-à-dire du CAP au BTS. Les niveaux 6 et 7 "
          "(licence, master) ne sont plus financés depuis 2025. Vos titres logistique sont donc "
          "dans la bonne fourchette : c'est un argument, utilisez-le."),
    ("T", "La condition n°1", "L'employeur doit s'être RECENSÉ auprès du CNFPT pendant la campagne "
          "annuelle (celle de 2026 était ouverte du 19 janvier au 20 mars 2026). "
          "Pas de recensement = pas de financement, même avec un bon candidat."),
    ("T", "La condition n°2", "Le métier visé doit figurer sur la liste des métiers en tension "
          "du CNFPT. VÉRIFIEZ que votre métier logistique y figure AVANT de promettre un "
          "financement à un DRH : c'est une promesse qui ne se rattrape pas."),
    ("T", "Les deux pièces à déposer", "① l'APF (accord préalable de financement) : déposé par la "
          "collectivité dans les 3 mois qui précèdent le début du contrat. "
          "② l'APC (demande de financement) : déposée par LE CFA dans le mois qui suit le "
          "démarrage. Cette deuxième pièce est de votre responsabilité — ne la laissez pas passer."),
    ("T", "À revérifier chaque année", "Les montants, les quotas et les dates de campagne changent "
          "d'une année sur l'autre. Ne citez jamais un chiffre sans l'avoir revérifié sur "
          "cnfpt.fr la semaine même."),
    ("H", "3. LA CONSÉQUENCE : UN CALENDRIER INVERSÉ", ""),
    ("T", "Le piège", "Dans le privé, on prospecte au printemps pour la rentrée. Dans le public "
          "territorial, si la collectivité ne s'est pas recensée en janvier-mars, la rentrée est "
          "déjà perdue pour le financement."),
    ("T", "Ce que ça change pour vous", "Vos rendez-vous d'octobre à décembre ne servent pas à "
          "signer : ils servent à faire inscrire la collectivité au recensement de janvier. "
          "Le closing, c'est « est-ce que je peux vous rappeler le 15 janvier pour votre "
          "recensement ? » — pas « est-ce que vous prenez un apprenti ? »."),
    ("T", "Et si la campagne est passée", "Deux options : demander au CNFPT s'il reste des contrats "
          "finançables sur l'année en cours, et surtout sécuriser l'intention pour la campagne "
          "suivante. Notez la date de rappel dans la colonne U de l'onglet concerné."),
    ("H", "4. LES DÉLAIS ADMINISTRATIFS À ANTICIPER", ""),
    ("T", "Ce qui prend du temps", "Inscription du poste au tableau des effectifs, délibération ou "
          "décision de l'exécutif, avis des instances, visa du comptable public. "
          "Comptez 3 mois entre l'accord oral et la signature — c'est normal, ce n'est pas un refus."),
    ("T", "Comment le gérer", "Faites du délai un argument : « je vous sollicite maintenant "
          "justement parce que je sais qu'il vous faut passer par une délibération »."),
    ("H", "5. LES OBJECTIONS SPÉCIFIQUES AU PUBLIC", ""),
    ("T", "« Nous n'avons pas de budget »",
          "« Les frais de formation ne sortent pas de votre budget : le CNFPT les prend en charge "
          "pour les niveaux CAP à BTS. Il vous reste la rémunération de l'apprenti, la plus basse "
          "du marché du travail. »"),
    ("T", "« Il faut une délibération »",
          "« Exactement, et c'est pour ça que je viens maintenant et pas en juillet. "
          "Je vous prépare la fiche de poste et le calendrier, vous n'avez plus qu'à la passer. »"),
    ("T", "« On ne pourra pas le garder après »",
          "Soyez honnête : l'apprentissage n'ouvre PAS droit à titularisation automatique — il faut "
          "le concours. Mais l'apprenti devient un candidat déjà formé à vos process, recrutable "
          "en contractuel, et prêt pour le concours. Un dispositif spécifique de titularisation "
          "existe pour les apprentis en situation de handicap : à vérifier au cas par cas."),
    ("T", "« Nous n'avons pas de tuteur »",
          "« Votre magasinier ou votre agent d'économat en est un. Je le forme et je "
          "l'accompagne, c'est compris dans notre accompagnement. »"),
    ("H", "6. VOS TROIS MEILLEURS ARGUMENTS DANS LE PUBLIC", ""),
    ("T", "① La transmission avant les départs",
          "Les magasiniers et agents d'économat partent en retraite et les savoirs partent avec eux. "
          "L'apprentissage est le seul moyen de transmettre avant le départ."),
    ("T", "② L'emploi des jeunes du territoire",
          "Un élu entend cet argument mieux que n'importe quel autre : vos candidats habitent la "
          "commune, et beaucoup n'ont pas de moyen de locomotion pour aller travailler ailleurs. "
          "À Sainte-Rose, Salazie ou La Plaine, la collectivité est parfois le seul employeur accessible."),
    ("T", "③ Le coût maîtrisé",
          "Formation prise en charge pour les niveaux CAP à BTS, rémunération réduite, "
          "aucun coût de recrutement. À condition d'avoir fait le recensement."),
]
r = 4
for typ, a, b in public:
    if typ == "H":
        c = ws.cell(row=r, column=1, value=a)
        c.font = F_H2; c.fill = FILL_JAUNE
        ws.cell(row=r, column=2).fill = FILL_JAUNE
        ws.row_dimensions[r].height = 20
    else:
        ca = ws.cell(row=r, column=1, value=a); ca.font = F_GRAS
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        cb = ws.cell(row=r, column=2, value=b); cb.font = F_CORPS
        cb.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = max(15, 14 * (len(b) // 103 + 1))
    r += 1
ws.sheet_view.showGridLines = False

# ---- Onglet 10 : script d'appel public
ws = wb.create_sheet("10 Script d'appel (public)")
entete_feuille(ws, "SCRIPT D'APPEL – COLLECTIVITÉS, ÉCOLES ET SANTÉ",
               "On ne parle pas à une mairie comme à un transporteur", 2)
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110
script = [
    ("H", "1. TROUVER LA BONNE PORTE", ""),
    ("T", "Commune / CCAS", "Demandez le DGS (directeur général des services) OU le directeur des "
          "services techniques. Formule : « je cherche la personne qui gère le magasin communal "
          "et les approvisionnements »."),
    ("T", "Lycée / collège", "Demandez l'ADJOINT GESTIONNAIRE. C'est lui qui tient le magasin, la "
          "demi-pension et les commandes. Ne demandez pas le proviseur en premier."),
    ("T", "Hôpital / CHU", "Demandez la direction des achats et de la logistique, puis la DRH. "
          "Le cadre du magasin général est celui qui dira « j'en ai besoin »."),
    ("T", "Clinique / EHPAD", "Demandez la DRH, le pharmacien gérant ou le responsable des services "
          "économiques. Dans un EHPAD, le directeur décide seul : c'est le circuit le plus court."),
    ("H", "2. L'ACCROCHE (30 SECONDES)", ""),
    ("T", "Phrase", "« Bonjour, [Prénom Nom], chargé des relations entreprises au CFA [nom]. "
          "Je ne vous appelle pas pour vous vendre une formation : je forme des jeunes de 16 à 25 ans "
          "aux métiers de la logistique sur le Nord-Est, et je cherche des structures qui ont un "
          "magasin ou un économat. Vous avez déjà accueilli un apprenti ? »"),
    ("T", "Pourquoi ça marche", "On annonce qu'on ne vend rien (essentiel dans le public, où l'on se "
          "méfie du démarchage), on nomme le territoire, et on termine par une question ouverte."),
    ("H", "3. LA QUESTION QUI QUALIFIE EN 30 SECONDES", ""),
    ("T", "À poser tôt", "« Qui réceptionne vos livraisons et qui tient votre stock aujourd'hui ? »"),
    ("T", "Ce que ça vous dit", "S'il y a une personne identifiée, il y a un poste ET un tuteur. "
          "Si la réponse est « tout le monde un peu », il y a un besoin d'organisation : c'est encore "
          "un meilleur argument."),
    ("H", "4. LE CLOSING SPÉCIFIQUE AU PUBLIC", ""),
    ("T", "Ne cherchez pas à signer", "Dans une collectivité, le premier appel ne se conclut jamais "
          "par une signature. Il se conclut par une DATE et un ENGAGEMENT DE RECENSEMENT."),
    ("T", "Phrase de closing", "« Je vous propose qu'on se voie 20 minutes. L'objectif : que vous "
          "soyez inscrit au recensement CNFPT de janvier, pour garder la porte ouverte. Ça ne vous "
          "engage à rien et ça vous coûte zéro. »"),
    ("T", "À noter systématiquement", "La date de rappel dans la colonne U, et l'état du recensement "
          "dans la colonne V (Notes). C'est votre indicateur n°1 de septembre à janvier."),
    ("H", "5. APRÈS L'APPEL – LE MÊME JOUR", ""),
    ("T", "Traçabilité", "Colonnes « Statut prospection », « Date dernier contact » et « Notes »."),
    ("T", "L'email qui vous démarque", "Envoyez sous 2h : la fiche de poste type de l'apprenti "
          "logistique, le calendrier CNFPT, et votre date de rappel. Un DRH de collectivité qui "
          "reçoit un calendrier propre vous rappelle."),
]
r = 4
for typ, a, b in script:
    if typ == "H":
        c = ws.cell(row=r, column=1, value=a)
        c.font = F_H2; c.fill = FILL_JAUNE
        ws.cell(row=r, column=2).fill = FILL_JAUNE
        ws.row_dimensions[r].height = 20
    else:
        ca = ws.cell(row=r, column=1, value=a); ca.font = F_GRAS
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        cb = ws.cell(row=r, column=2, value=b); cb.font = F_CORPS
        cb.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = max(15, 14 * (len(b) // 105 + 1))
    r += 1
ws.sheet_view.showGridLines = False

# ---- Onglet 11 : sources
ws = wb.create_sheet("11 Sources & mise à jour")
entete_feuille(ws, "SOURCES DES DONNÉES",
               "Traçabilité et outils pour compléter / actualiser le fichier", 2)
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 92
src = [
    ("Annuaire des Entreprises (INSEE) – y compris collectivités", "https://annuaire-entreprises.data.gouv.fr/"),
    ("CNFPT – financement des contrats d'apprentissage", "https://www.cnfpt.fr/s-informer/nos-actualites/le-fil-dactu/financement-contrats-dapprentissage/national"),
    ("CNFPT – accueillir un apprenti dans une collectivité", "https://www.cnfpt.fr/se-former/accueillir-apprenti/lapprentissage-collectivites-territoriales"),
    ("Annuaire de l'administration (Service-Public)", "https://lannuaire.service-public.gouv.fr/la-reunion/la-reunion"),
    ("Annuaire de l'éducation (établissements scolaires)", "https://data.education.gouv.fr/explore/dataset/fr-en-annuaire-education/"),
    ("Santé.fr – annuaire des établissements de santé (FINESS)", "https://www.sante.fr/"),
    ("CIREST – communauté intercommunale Réunion Est", "https://www.cirest.fr/qui-sommes-nous/"),
    ("CHU de La Réunion", "https://www.chu-reunion.fr/contact/"),
    ("Groupe de santé Clinifutur", "https://www.clinifutur.net/"),
    ("AURAR – centres de dialyse", "https://aurar.fr/nous-trouver/"),
]
r = 4
ws.cell(row=r, column=1, value="SOURCES UTILISÉES POUR CONSTRUIRE CE FICHIER").font = F_H2
r += 1
for lib, url in src:
    ws.cell(row=r, column=1, value=lib).font = F_CORPS
    c = ws.cell(row=r, column=2, value=url)
    c.font = F_LIEN; c.hyperlink = url
    r += 1
r += 1
ws.cell(row=r, column=1, value="LIMITES CONNUES").font = F_H2
r += 1
for txt in [
    "• Les NOMS des structures sont publics et fiables. Les ADRESSES proviennent de sources publiques et doivent être confirmées : la colonne « Fiabilité donnée » indique le niveau de confiance de chaque ligne.",
    "• Les règles de financement CNFPT (montants, quotas, dates de campagne, liste des métiers en tension) changent chaque année. Revérifiez sur cnfpt.fr avant d'annoncer quoi que ce soit à un employeur.",
    "• Les lignes « à lister » (autres lycées et collèges, EHPAD, ESAT, autres cliniques du groupe) sont des méthodes de recherche, pas des structures : à vous de les déplier commune par commune.",
    "• Le GHER (Saint-Benoît) figure aussi dans le fichier de prospection logistique privée : un seul appel, pas deux.",
    "• L'EPSMR est situé à Saint-Paul, hors du bassin Nord-Est : il n'est là que pour un candidat mobile.",
]:
    c = ws.cell(row=r, column=1, value=txt); c.font = F_CORPS
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.row_dimensions[r].height = max(15, 14 * (len(txt) // 130 + 1))
    r += 1
ws.sheet_view.showGridLines = False

out = "/home/user/Claude-/prospection/Prospection_Public_Sante_Nord-Est_Reunion.xlsx"
wb.save(out)
print("OK ->", out)
print("Onglets :", wb.sheetnames)
print("Total structures :", sum(n for _, n in compteurs))
