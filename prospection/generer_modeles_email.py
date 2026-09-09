# -*- coding: utf-8 -*-
"""
Genere un classeur Excel lisible listant tous les modeles d'email de prospection
(entreprises privees, secteur public, candidats), a partir de modeles_email.py.

Ce classeur sert de reference / de brouillon a modifier a la main. Le texte qui
compte reellement pour la fusion automatique reste celui de modeles_email.py :
si vous modifiez un objet ou un message ici dans Excel, reportez le changement
dans modeles_email.py pour qu'il soit pris en compte par generer_emails_personnalises.py.

Usage :
    python3 generer_modeles_email.py
"""
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl n'est pas installe.\n"
             "Lancez d'abord :  python3 -m pip install openpyxl")

from commun_prospection import (
    F_TITRE, F_STITRE, F_ENTETE, F_CORPS, F_GRAS, F_H2, F_LIEN,
    FILL_ENTETE, FILL_ALT, FILL_JAUNE, BORD, entete_feuille,
)
from modeles_email import CONFIG, TEMPLATES, variables_disponibles

TITRES_CATEGORIES = {
    "entreprise_prive": ("01 Entreprises privées", "MODÈLES — ENTREPRISES PRIVÉES",
                          "A utiliser sur le fichier Prospection_Logistique_Nord-Est_Reunion.xlsx"),
    "secteur_public": ("02 Secteur public", "MODÈLES — SECTEUR PUBLIC ET SANTÉ",
                        "A utiliser sur le fichier Prospection_Public_Sante_Nord-Est_Reunion.xlsx"),
    "candidats": ("03 Candidats", "MODÈLES — CANDIDATS (16-25 ANS)",
                  "Réponses aux candidats, hors fusion automatique (pas de fichier source dédié)"),
}

COLS_MODELE = [("Clé", 26), ("Nom du modèle", 30), ("Quand l'utiliser", 42),
               ("Objet", 42), ("Message", 90)]


def onglet_modeles(wb, categorie):
    nom_onglet, titre, sous_titre = TITRES_CATEGORIES[categorie]
    ws = wb.create_sheet(nom_onglet)
    entete_feuille(ws, titre, sous_titre, len(COLS_MODELE))

    r = 4
    for i, (lib, larg) in enumerate(COLS_MODELE, start=1):
        c = ws.cell(row=r, column=i, value=lib)
        c.font, c.fill, c.border = F_ENTETE, FILL_ENTETE, BORD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[r].height = 34

    for j, (cle, modele) in enumerate(TEMPLATES[categorie].items(), start=1):
        row = r + j
        vals = [cle, modele["nom"], modele["quand"], modele["sujet"], modele["corps"]]
        for k, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=k, value=v)
            c.font, c.border = F_CORPS, BORD
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if j % 2 == 0:
                c.fill = FILL_ALT
        ws.cell(row=row, column=1).font = F_GRAS
        ws.cell(row=row, column=2).font = F_GRAS
        nb_lignes = max(modele["corps"].count("\n") + 1, 6)
        ws.row_dimensions[row].height = min(15 * nb_lignes, 400)

    ws.auto_filter.ref = f"A{r}:{get_column_letter(len(COLS_MODELE))}{r + len(TEMPLATES[categorie])}"
    ws.freeze_panes = f"A{r + 1}"
    ws.sheet_view.zoomScale = 90
    return ws


def onglet_variables(wb):
    ws = wb.create_sheet("04 Variables disponibles")
    entete_feuille(ws, "VARIABLES DISPONIBLES DANS LES MODÈLES",
                    "Ce que chaque {variable} devient automatiquement dans l'email généré", 2)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 80

    r = 4
    ws.cell(row=r, column=1, value="Variable").font = F_ENTETE
    ws.cell(row=r, column=2, value="Ce qu'elle devient").font = F_ENTETE
    for k in (1, 2):
        ws.cell(row=r, column=k).fill = FILL_ENTETE
        ws.cell(row=r, column=k).border = BORD
    r += 1
    for var, desc in variables_disponibles():
        ws.cell(row=r, column=1, value=var).font = F_GRAS
        ws.cell(row=r, column=2, value=desc).font = F_CORPS
        for k in (1, 2):
            ws.cell(row=r, column=k).border = BORD
            ws.cell(row=r, column=k).alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="VOS RÉGLAGES ACTUELS (CONFIG dans modeles_email.py)").font = F_H2
    r += 1
    for k, v in CONFIG.items():
        ws.cell(row=r, column=1, value="{%s}" % k).font = F_GRAS
        ws.cell(row=r, column=2, value=v).font = F_CORPS
        r += 1
    ws.cell(row=r + 1, column=1,
            value="Pour changer ces valeurs, modifiez la section CONFIG en haut du fichier "
                  "modeles_email.py, puis relancez ce script.").font = F_STITRE
    ws.sheet_view.showGridLines = False


def onglet_mode_emploi(wb):
    ws = wb.create_sheet("00 Mode d'emploi")
    entete_feuille(ws, "SYSTÈME D'EMAIL DE PROSPECTION", "Comment ça marche, en 3 étapes", 1)
    ws.column_dimensions["A"].width = 110
    lignes = [
        ("", None),
        ("1. Personnalisez vos coordonnées", F_H2),
        ("Ouvrez modeles_email.py et complétez la section CONFIG en haut du fichier "
         "(nom du CFA, votre nom, téléphone, email). C'est fait une seule fois.", F_CORPS),
        ("", None),
        ("2. Ce classeur : la référence à lire", F_H2),
        ("Les onglets 01 à 03 listent tous les modèles disponibles, avec l'objet, le message "
         "complet et le moment où les utiliser. Relisez-les avant le premier envoi, adaptez le "
         "ton si besoin (le texte qui compte pour la génération automatique reste celui de "
         "modeles_email.py).", F_CORPS),
        ("", None),
        ("3. Générer les emails personnalisés", F_H2),
        ("Le script generer_emails_personnalises.py lit vos fichiers de prospection "
         "(Prospection_Logistique_Nord-Est_Reunion.xlsx ou "
         "Prospection_Public_Sante_Nord-Est_Reunion.xlsx), prend chaque ligne où un email est "
         "connu, et génère un brouillon d'email personnalisé par entreprise.", F_CORPS),
        ("Exemple :", F_GRAS),
        ("   python3 generer_emails_personnalises.py --categorie entreprise_prive "
         "--modele PREMIER_CONTACT --statut \"A contacter\"", F_CORPS),
        ("Le script écrit un fichier .eml par entreprise dans un dossier emails_generes/. "
         "Un .eml s'ouvre d'un double-clic dans Outlook, Mail ou Thunderbird : le message "
         "s'affiche déjà rempli (destinataire, objet, texte), prêt à être relu puis envoyé "
         "à la main. RIEN N'EST ENVOYÉ AUTOMATIQUEMENT.", F_GRAS),
        ("Un fichier fusion_emails.csv est produit en complément, pour les outils de "
         "publipostage (Gmail, Outlook) si vous préférez cette méthode.", F_CORPS),
        ("", None),
        ("Rappel RGPD / CNIL", F_H2),
        ("La prospection professionnelle par email vers une adresse professionnelle est "
         "autorisée sans accord préalable dès lors qu'elle concerne l'activité du "
         "destinataire, à condition d'identifier clairement l'expéditeur et d'offrir un "
         "moyen simple de s'y opposer : c'est le rôle de la mention « STOP » incluse dans "
         "la signature des emails entreprises et secteur public.", F_CORPS),
    ]
    r = 4
    for texte, font in lignes:
        c = ws.cell(row=r, column=1, value=texte)
        if font:
            c.font = font
        c.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    ws.sheet_view.showGridLines = False


def main():
    wb = Workbook()
    wb.remove(wb.active)
    onglet_mode_emploi(wb)
    for categorie in ("entreprise_prive", "secteur_public", "candidats"):
        onglet_modeles(wb, categorie)
    onglet_variables(wb)

    ici = os.path.dirname(os.path.abspath(__file__))
    dest = os.path.join(ici, "Modeles_Email_Prospection.xlsx")
    wb.save(dest)
    print("Classeur généré : %s" % dest)


if __name__ == "__main__":
    main()
